from __future__ import annotations
from pathlib import Path
from typing import Any
import csv, json


def load_tabular(path: str | Path, sheet: str | None = None) -> list[dict[str, Any]]:
    p=Path(path)
    ext=p.suffix.lower()
    if ext=='.csv':
        with p.open('r',encoding='utf-8-sig',newline='') as f:
            return list(csv.DictReader(f))
    if ext in {'.json','.jsonl'}:
        if ext=='.jsonl':
            return [json.loads(line) for line in p.read_text(encoding='utf-8').splitlines() if line.strip()]
        data=json.loads(p.read_text(encoding='utf-8'))
        if isinstance(data,list): return data
        if isinstance(data,dict):
            for key in ['rows','items','data','records']:
                if isinstance(data.get(key),list): return data[key]
        raise ValueError('JSON does not contain a row list')
    if ext in {'.xlsx','.xlsm'}:
        from openpyxl import load_workbook
        wb=load_workbook(p,read_only=True,data_only=True)
        ws=wb[sheet] if sheet else wb[wb.sheetnames[0]]
        rows=ws.iter_rows(values_only=True)
        headers=[str(x).strip() if x is not None else '' for x in next(rows)]
        out=[]
        for vals in rows:
            row={headers[i]:vals[i] for i in range(min(len(headers),len(vals))) if headers[i]}
            if any(v not in (None,'') for v in row.values()): out.append(row)
        return out
    raise ValueError(f'unsupported tabular file: {ext}')


def apply_field_map(rows:list[dict[str,Any]], field_map:dict[str,str]|None=None)->list[dict[str,Any]]:
    if not field_map: return rows
    return [{canonical:r.get(source) for canonical,source in field_map.items()} | {k:v for k,v in r.items() if k not in field_map.values()} for r in rows]
